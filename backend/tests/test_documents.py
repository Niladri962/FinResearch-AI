"""Parsing, table understanding, metadata extraction and chunking."""
from __future__ import annotations

from pathlib import Path

import pytest

from app.documents.metadata import classify_document, company_key, detect_company, extract_metadata
from app.documents.parser import parse_document
from app.documents.table_extractor import (
    detect_currency,
    detect_text_tables,
    detect_unit,
    extract_structured_table,
    is_header_line,
    parse_number,
    split_row,
    table_to_markdown,
)
from app.rag.chunking import FinancialChunker, embedding_text
from app.utils.errors import EmptyDocumentError, InvalidDocumentError, UnsupportedFormatError

AR25 = "aurora_industries_annual_report_fy2025.pdf"
TRANSCRIPT = "aurora_industries_q4_fy2025_earnings_call_transcript.pdf"
BOREALIS = "borealis_technologies_annual_report_fy2025.pdf"


# ── Numbers, units, rows ─────────────────────────────────────────────────────
class TestNumberParsing:
    @pytest.mark.parametrize(
        ("raw", "expected"),
        [
            ("1,234.50", 1234.5), ("(1,730)", -1730.0), ("-45", -45.0), ("₹ 2,500", 2500.0),
            ("1,23,456", 123456.0), ("$3.30", 3.3), ("0.5", 0.5),
        ],
    )
    def test_values(self, raw, expected):
        assert parse_number(raw) == expected

    @pytest.mark.parametrize("raw", ["", "-", "nil", "12.5%", "N/A", "abc", "12a", None])
    def test_non_values(self, raw):
        assert parse_number(raw) is None

    def test_units_and_currency(self):
        assert detect_unit("(INR in crore, except per share data)") == ("crore", 1e7)
        assert detect_unit("Amounts in USD millions") == ("million", 1e6)
        assert detect_unit("₹ in lakhs") == ("lakh", 1e5)
        assert detect_unit("no unit here") is None
        assert detect_currency("(INR in crore)") == "INR"
        assert detect_currency("USD in millions") == "USD"
        assert detect_currency("Revenue grew strongly") is None

    def test_split_row(self):
        assert split_row("Revenue from operations   10,600   9,400") == ("Revenue from operations", ["10,600", "9,400"])
        assert split_row("Finance costs 21 (240) (170)") == ("Finance costs", ["21", "(240)", "(170)"])
        assert split_row("This sentence has no figures at the end.") is None
        assert split_row("12,500 10,000") is None

    def test_header_line(self):
        assert is_header_line("Particulars   FY2025   FY2024")
        assert is_header_line("Year ended March 31, 2025   Year ended March 31, 2024")
        assert not is_header_line("Revenue grew strongly in FY2025 compared with FY2024")


# ── Table understanding ──────────────────────────────────────────────────────
class TestTableExtraction:
    def test_structured_facts_with_units(self):
        rows = [
            ["Particulars", "FY2025", "FY2024"],
            ["Revenue from operations", "1,250", "1,000"],
            ["Finance costs", "(40)", "(30)"],
            ["Net profit margin", "12.0%", "11.0%"],
            ["Some unrecognised line", "5", "4"],
            ["Profit for the year", "150", "110"],
        ]
        table = extract_structured_table(rows, title="Statement of Profit and Loss (INR in crore)")
        assert table.periods == ["FY2025", "FY2024"]
        assert (table.currency, table.unit, table.scale) == ("INR", "crore", 1e7)
        facts = {(f.metric, f.period.label): f.value for f in table.facts}
        assert facts[("revenue", "FY2025")] == 1250.0
        assert facts[("revenue", "FY2024")] == 1000.0
        assert facts[("finance_costs", "FY2025")] == -40.0
        assert facts[("net_income", "FY2024")] == 110.0
        assert not any(metric.endswith("margin") for metric, _ in facts)       # percentage rows are skipped
        # The spec's structured representation: one record per row, keyed by period.
        record = table.as_records()[0]
        assert record == {"metric": "revenue", "label": "Revenue from operations", "FY2025": 1250.0,
                          "FY2024": 1000.0, "currency": "INR", "unit": "crore"}

    def test_no_period_header_means_no_facts(self):
        rows = [["Item", "Amount"], ["Revenue", "100"], ["Net profit", "10"]]
        assert extract_structured_table(rows) is None

    def test_note_column_is_ignored(self):
        rows = [["Particulars", "Note", "FY2025", "FY2024"], ["Revenue from operations", "21", "500", "400"]]
        table = extract_structured_table(rows)
        assert {(f.period.label, f.value) for f in table.facts} == {("FY2025", 500.0), ("FY2024", 400.0)}

    def test_borrowings_resolved_by_balance_sheet_section(self):
        rows = [
            ["Particulars", "As at March 31, 2025", "As at March 31, 2024"],
            ["Non-current liabilities"],
            ["Borrowings", "2,600", "1,700"],
            ["Current liabilities"],
            ["Borrowings", "900", "600"],
            ["Total equity and liabilities", "11,200", "9,400"],
        ]
        facts = {(f.metric, f.period.label): f.value for f in extract_structured_table(rows).facts}
        assert facts[("long_term_debt", "FY2025")] == 2600.0
        assert facts[("short_term_debt", "FY2025")] == 900.0
        assert not any(metric in ("equity", "total_liabilities") for metric, _ in facts)

    def test_eps_only_under_earnings_per_share_heading(self):
        rows = [["Particulars", "FY2025", "FY2024"], ["Earnings per equity share"], ["Basic (INR)", "14.40", "16.32"]]
        facts = extract_structured_table(rows).facts
        assert {(f.metric, f.value) for f in facts} == {("eps", 14.4), ("eps", 16.32)}
        orphan = [["Particulars", "FY2025", "FY2024"], ["Basic", "14.40", "16.32"]]
        table = extract_structured_table(orphan)
        assert table is None or table.facts == []

    def test_interim_columns_are_not_treated_as_annual(self):
        rows = [["Particulars", "Quarter ended June 30, 2025", "Quarter ended June 30, 2024"], ["Revenue", "300", "250"]]
        kinds = {f.period.kind for f in extract_structured_table(rows).facts}
        assert kinds == {"interim"}

    def test_detect_text_tables(self):
        lines = [
            "Management commentary paragraph that should stay as prose.",
            "Particulars   FY2025   FY2024",
            "ASSETS",
            "Inventories   1,620   1,300",
            "Trade receivables   1,420   1,180",
            "Total assets   11,200   9,400",
            "The notes form an integral part of these statements.",
        ]
        tables = detect_text_tables(lines)
        assert len(tables) == 1
        assert (tables[0].start, tables[0].end) == (1, 6)
        assert tables[0].rows[0] == ["Particulars", "FY2025", "FY2024"]
        assert tables[0].rows[-1] == ["Total assets", "11,200", "9,400"]

    def test_prose_with_numbers_is_not_a_table(self):
        lines = ["Revenue grew to 10,600", "which was above our plan.", "Costs rose too."]
        assert detect_text_tables(lines) == []

    def test_markdown_rendering(self):
        assert table_to_markdown([["A", "B"], ["x", "1"]]) == "| A | B |\n| --- | --- |\n| x | 1 |"


# ── Parsers ──────────────────────────────────────────────────────────────────
class TestParsers:
    def test_pdf_layout_elements(self, sample_docs):
        parsed = parse_document(sample_docs[AR25])
        assert parsed.page_count == 6
        headings = [e.text for page in parsed.pages for e in page.elements if e.kind == "heading"]
        assert {"Risk Factors", "Management Discussion and Analysis", "Consolidated Balance Sheet"} <= set(headings)
        tables = [e for page in parsed.pages for e in page.elements if e.kind == "table"]
        assert len(tables) == 4                      # highlights, P&L, balance sheet (unruled), cash flow
        # Running headers/footers and page numbers are stripped.
        assert "SYNTHETIC SAMPLE - fictitious" not in parsed.pages[2].text
        assert "| Annual Report FY2025" not in parsed.pages[2].text

    def test_pdf_ruled_and_unruled_tables_are_both_recovered(self, sample_docs):
        parsed = parse_document(sample_docs[AR25])
        balance_sheet = next(e for e in parsed.pages[4].elements if e.kind == "table")
        labels = [row[0] for row in balance_sheet.rows]
        assert "Total current liabilities" in labels and "Borrowings" in labels
        assert ["Total assets", "11,200", "9,400"] in balance_sheet.rows

    def test_text_file_with_headings_and_table(self, tmp_path: Path):
        path = tmp_path / "notes.md"
        path.write_text(
            "# Acme Widgets Limited\n\nAnnual Report FY2025\n\n## Risk Factors\n\n"
            "Demand for widgets is cyclical and a downturn would adversely affect revenue and margins.\n\n"
            "## Financial Summary\n\n(INR in crore)\n\n"
            "| Particulars | FY2025 | FY2024 |\n|---|---|---|\n| Revenue | 500 | 400 |\n"
            "| Net profit | 50 | 30 |\n| Total assets | 900 | 800 |\n",
            encoding="utf-8",
        )
        parsed = parse_document(path)
        kinds = [e.kind for e in parsed.pages[0].elements]
        assert kinds.count("table") == 1 and "heading" in kinds
        table = next(e for e in parsed.pages[0].elements if e.kind == "table")
        assert ["Revenue", "500", "400"] in table.rows

    def test_docx(self, tmp_path: Path):
        from docx import Document

        doc = Document()
        doc.add_heading("Globex Corporation", level=1)
        doc.add_paragraph("Quarterly report for the quarter ended June 30, 2025. Amounts in USD millions.")
        doc.add_heading("Results of Operations", level=2)
        doc.add_paragraph("Revenue increased on higher subscription volumes and improved pricing across regions.")
        table = doc.add_table(rows=3, cols=3)
        for r, row in enumerate([["Particulars", "FY2025", "FY2024"], ["Revenue", "120", "100"], ["Net income", "12", "9"]]):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        path = tmp_path / "globex.docx"
        doc.save(str(path))

        parsed = parse_document(path)
        elements = parsed.pages[0].elements
        assert [e.kind for e in elements] == ["heading", "text", "heading", "text", "table"]
        assert elements[-1].rows[1] == ["Revenue", "120", "100"]

    def test_xlsx(self, tmp_path: Path):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Income Statement"
        for row in [["Particulars", 2025, 2024], ["Revenue", 1250.5, 1000], ["Net income", 150, 110]]:
            sheet.append(row)
        path = tmp_path / "statements.xlsx"
        workbook.save(str(path))

        parsed = parse_document(path)
        table = next(e for e in parsed.pages[0].elements if e.kind == "table")
        assert table.rows[0] == ["Particulars", "2025", "2024"]
        facts = {(f.metric, f.period.label): f.value for f in extract_structured_table(table.rows).facts}
        assert facts[("revenue", "FY2025")] == 1250.5 and facts[("net_income", "FY2024")] == 110.0

    def test_invalid_pdf(self, tmp_path: Path):
        path = tmp_path / "broken.pdf"
        path.write_bytes(b"%PDF-1.7 this is not really a pdf")
        with pytest.raises(InvalidDocumentError):
            parse_document(path)

    def test_empty_text(self, tmp_path: Path):
        path = tmp_path / "empty.txt"
        path.write_text("   \n", encoding="utf-8")
        with pytest.raises(EmptyDocumentError):
            parse_document(path)

    def test_pdf_without_text(self, tmp_path: Path):
        import pymupdf

        doc = pymupdf.open()
        doc.new_page()
        path = tmp_path / "blank.pdf"
        doc.save(str(path))
        doc.close()
        with pytest.raises(EmptyDocumentError):
            parse_document(path)

    def test_unsupported_extension(self, tmp_path: Path):
        path = tmp_path / "data.csv"
        path.write_text("a,b\n1,2\n", encoding="utf-8")
        with pytest.raises(UnsupportedFormatError):
            parse_document(path)


# ── Metadata ─────────────────────────────────────────────────────────────────
class TestMetadata:
    def test_annual_report(self, sample_docs):
        meta = extract_metadata(parse_document(sample_docs[AR25]), AR25)
        assert meta.company == "Aurora Industries Limited"
        assert meta.document_type == "annual_report"
        assert (meta.fiscal_year, meta.quarter, meta.reporting_period) == (2025, None, "FY2025")
        assert meta.currency == "INR" and meta.unit == ("crore", 1e7)
        assert meta.title == "Annual Report FY2025"

    def test_transcript(self, sample_docs):
        meta = extract_metadata(parse_document(sample_docs[TRANSCRIPT]), TRANSCRIPT)
        assert meta.document_type == "earnings_call_transcript"
        assert (meta.fiscal_year, meta.quarter, meta.reporting_period) == (2025, 4, "Q4 FY2025")

    def test_second_company_and_currency(self, sample_docs):
        meta = extract_metadata(parse_document(sample_docs[BOREALIS]), BOREALIS)
        assert meta.company == "Borealis Technologies Inc."
        assert meta.currency == "USD" and meta.unit == ("million", 1e6)

    def test_user_overrides_win(self, sample_docs):
        meta = extract_metadata(
            parse_document(sample_docs[AR25]), AR25,
            {"company": "Aurora Group", "fiscal_year": 2030, "quarter": 2, "document_type": "quarterly_report"},
        )
        assert meta.company == "Aurora Group"
        assert (meta.document_type, meta.reporting_period) == ("quarterly_report", "Q2 FY2030")

    @pytest.mark.parametrize(
        ("text", "filename", "expected"),
        [
            ("Form 10-K annual report for the fiscal year", "acme_10k.pdf", "annual_report"),
            ("Operator: welcome to the earnings call. Question-and-answer session follows.", "call.txt", "earnings_call_transcript"),
            ("Investor Presentation. Safe harbor statement.", "deck.pdf", "investor_presentation"),
            ("Initiating coverage with a target price of 120. Equity research.", "note.pdf", "research_report"),
            ("Nothing recognisable in here at all.", "scan001.pdf", "other"),
        ],
    )
    def test_classification(self, text, filename, expected):
        assert classify_document(text, filename) == expected

    def test_company_helpers(self):
        assert company_key("ACME Widgets Ltd.") == company_key("Acme Widgets Limited") == "acme widgets"
        assert detect_company("Welcome to Globex Corporation. Globex Corporation is a leader.", "x.pdf") == "Globex Corporation"
        assert detect_company("no company here", "initech_annual_report_fy2024.pdf") == "Initech"


# ── Chunking ─────────────────────────────────────────────────────────────────
class TestChunking:
    @pytest.fixture(scope="class")
    def chunks(self, sample_docs):
        parsed = parse_document(sample_docs[AR25])
        return FinancialChunker(target_tokens=320, overlap_tokens=48).chunk(parsed, "annual_report")

    def test_tables_are_single_intact_chunks(self, chunks):
        tables = [c for c in chunks if c.table_rows]
        assert len(tables) == 4
        balance = next(c for c in tables if "Balance Sheet" in c.section)
        assert balance.chunk_type == "financial_statement"
        # Header, caption (with the unit) and every row travel together.
        assert "(INR in crore)" in balance.text
        assert "| Particulars | FY2025 | FY2024 |" in balance.text
        assert "| Total equity and liabilities | 11,200 | 9,400 |" in balance.text

    def test_every_chunk_has_source_metadata(self, chunks):
        for chunk in chunks:
            assert chunk.page_start >= 1 and chunk.page_end >= chunk.page_start
            assert chunk.chunk_type and chunk.token_count > 0
        assert all(c.section for c in chunks)

    def test_semantic_chunk_types(self, chunks):
        by_section = {c.section: c.chunk_type for c in chunks}
        assert by_section["Risk Factors"] == "risk_factor"
        assert by_section["Significant Accounting Policies"] == "accounting_policy"
        assert by_section["Management Discussion and Analysis › Outlook"] == "management_commentary"
        assert by_section["Chairman's Letter"] == "management_commentary"

    def test_sections_are_not_mixed(self, chunks):
        risk = next(c for c in chunks if c.section == "Risk Factors")
        assert "Customer concentration" in risk.text and "Chairman" not in risk.text

    def test_no_page_furniture_chunks(self, chunks):
        assert all(c.token_count >= 12 for c in chunks)

    def test_transcript_chunks_are_earnings_commentary(self, sample_docs):
        parsed = parse_document(sample_docs[TRANSCRIPT])
        chunks = FinancialChunker().chunk(parsed, "earnings_call_transcript")
        assert chunks and {c.chunk_type for c in chunks} == {"earnings_commentary"}

    def test_long_narrative_is_split_on_sentences_with_overlap(self):
        from app.documents.parser import Element, ParsedDocument, ParsedPage

        sentences = [f"Sentence number {i} explains one aspect of the business in plain words." for i in range(60)]
        page = ParsedPage(number=1, elements=[
            Element(kind="heading", text="Business Overview"),
            *[Element(kind="text", text=" ".join(sentences[i:i + 6])) for i in range(0, 60, 6)],
        ])
        chunks = FinancialChunker(target_tokens=120, overlap_tokens=30).chunk(ParsedDocument("text", [page]))
        assert len(chunks) > 3
        assert all(c.token_count <= 120 + 30 + 20 for c in chunks)
        assert all(c.text.rstrip().endswith(".") for c in chunks)          # never cut mid-sentence
        assert chunks[1].text.split(". ")[0].startswith("Sentence number")  # overlap carries whole sentences
        first_tail = chunks[0].text.rstrip().split(". ")[-1]
        assert first_tail.rstrip(".") in chunks[1].text

    def test_oversized_table_is_split_with_repeated_header(self):
        from app.documents.parser import Element, ParsedDocument, ParsedPage

        rows = [["Particulars", "FY2025", "FY2024"]] + [[f"Line item {i}", f"{i * 10:,}", f"{i * 9:,}"] for i in range(1, 121)]
        page = ParsedPage(number=3, elements=[Element(kind="heading", text="Schedule of Items"), Element(kind="table", rows=rows)])
        chunks = FinancialChunker(max_table_tokens=300).chunk(ParsedDocument("text", [page]))
        assert len(chunks) > 1
        assert all("| Particulars | FY2025 | FY2024 |" in c.text for c in chunks)
        assert chunks[0].table_rows is not None and all(c.table_rows is None for c in chunks[1:])
        joined = "\n".join(c.text for c in chunks)
        assert all(f"| Line item {i} |" in joined for i in (1, 60, 120))

    def test_embedding_text_carries_provenance(self):
        text = embedding_text("Body.", company="Acme", title="Annual Report FY2025", section="Risk Factors")
        assert text == "Acme | Annual Report FY2025 | Risk Factors\nBody."
