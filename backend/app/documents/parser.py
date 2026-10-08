"""Document parsing.

Turns PDF / DOCX / XLSX / text files into an ordered stream of layout elements
(headings, paragraphs, tables) per page. Layout is preserved because the
chunker relies on it: tables stay whole and headings define sections.

PDF strategy
  * PyMuPDF supplies text spans with position, font size and weight.
  * pdfplumber finds ruled tables; text inside their bounding boxes is taken
    from the table grid rather than the text layer, so nothing is duplicated.
  * Unruled tables are recovered from aligned text rows.
  * Pages without a text layer go through OCR when Tesseract is available.
"""
from __future__ import annotations

import re
import statistics
import time
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

from app.documents.table_extractor import detect_text_tables, is_header_fragment, is_numeric_cell, split_row
from app.utils.errors import EmptyDocumentError, InvalidDocumentError, OCRError, UnsupportedFormatError
from app.utils.logging import get_logger

logger = get_logger(__name__)

MAX_PAGES = 2000
RULED_TABLE_BUDGET_SECONDS = 20.0
_BULLET = re.compile(r"^\s*(?:[•▪■●◦‣·*]|[-–—]\s|\(?\d{1,2}[.)]\s|\(?[a-z][.)]\s)")
_PAGE_NUMBER = re.compile(r"(?i)^(page\s*)?\d{1,4}(\s*(of|/|\|)\s*\d{1,4})?$")
KNOWN_SECTION = re.compile(
    r"(?i)^(?:\d{1,2}[.)]?\s+|item\s+\d+[a-z]?[.:]?\s+)?("
    r"risk factors?|risk management|key risks?|principal risks.*|"
    r"management'?s? discussion and analysis.*|md&a|business overview|company overview|"
    r"chairman'?s? (letter|message|statement)|(ceo|md|cfo)'?s? (letter|message|review)|"
    r"letter to shareholders|outlook|future outlook|guidance|strategy|strategic priorities|"
    r"financial (highlights|performance|review|statements|summary)|results of operations|"
    r"(consolidated |standalone )?(balance sheet|statement of (profit and loss|cash flows?|financial position|"
    r"comprehensive income|changes in equity)|income statement|cash flow statement)|"
    r"notes to (the )?(consolidated |standalone )?financial statements|"
    r"(significant|material|summary of) accounting policies|"
    r"liquidity and capital resources|capital expenditure|segment (information|performance|reporting)|"
    r"corporate governance|directors'? report|auditors?'? report|independent auditors?'? report|"
    r"question[- ]and[- ]answer( session)?|q\s*&\s*a( session)?|prepared remarks|opening remarks|"
    r"opportunities and threats|growth opportunities|revenue|profitability"
    r")\s*$"
)


@dataclass
class Element:
    kind: Literal["heading", "text", "table"]
    text: str = ""
    rows: list[list[str]] | None = None
    level: int = 2
    y: float = 0.0


@dataclass
class ParsedPage:
    number: int
    elements: list[Element] = field(default_factory=list)

    @property
    def text(self) -> str:
        parts = []
        for el in self.elements:
            if el.kind == "table" and el.rows:
                parts.extend("  ".join(c for c in row if c) for row in el.rows)
            else:
                parts.append(el.text)
        return "\n".join(parts)


@dataclass
class ParsedDocument:
    file_type: str
    pages: list[ParsedPage] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def page_count(self) -> int:
        return len(self.pages)

    def text(self, max_pages: int | None = None) -> str:
        return "\n".join(p.text for p in self.pages[:max_pages])

    @property
    def char_count(self) -> int:
        return sum(len(p.text) for p in self.pages)


@dataclass
class Line:
    text: str
    y0: float
    y1: float
    size: float = 10.0
    bold: bool = False
    forced_heading: int = 0  # >0 when the source format marks it as a heading (level)


# ── Shared layout logic ──────────────────────────────────────────────────────
def is_heading(line: Line, body_size: float) -> int:
    """Return heading level (1 or 2) or 0."""
    if line.forced_heading:
        return line.forced_heading
    text = line.text.strip()
    words = text.split()
    if len(text) < 3 or len(text) > 120 or len(words) > 14:
        return 0
    if text[-1] in ".;,:":
        return 0
    known = bool(KNOWN_SECTION.match(text))
    styled = line.size >= body_size * 1.15 or line.bold
    if _BULLET.match(text) and not (known or styled):
        return 0  # an unstyled list item, not a numbered heading
    if split_row(text) is not None or sum(is_numeric_cell(w) for w in words) > len(words) / 2:
        return 0
    if is_header_fragment(text):
        return 0  # bold column headers ("Particulars   Note") are not section headings
    if not re.search(r"[A-Za-z]{3}", text):
        return 0
    if line.size >= body_size * 1.4:
        return 1
    if line.size >= body_size * 1.15:
        return 2
    if line.bold and len(words) <= 10:
        return 2
    if text.isupper() and len(words) <= 8:
        return 2
    if KNOWN_SECTION.match(text):
        return 2
    return 0


def _join_lines(lines: list[str]) -> str:
    out = ""
    for raw in lines:
        text = re.sub(r"\s+", " ", raw).strip()
        if not text:
            continue
        if out.endswith("-") and text[:1].islower() and len(out) > 1 and out[-2].isalpha():
            out = out[:-1] + text  # re-join a word hyphenated across lines
        else:
            out = f"{out} {text}" if out else text
    return out


def build_elements(lines: list[Line], body_size: float) -> list[Element]:
    """Classify a page's lines into headings, paragraphs and text tables."""
    elements: list[Element] = []
    table_at = {t.start: t for t in detect_text_tables([ln.text for ln in lines])}
    paragraph: list[Line] = []

    def flush() -> None:
        if paragraph:
            text = _join_lines([ln.text for ln in paragraph])
            if text:
                elements.append(Element(kind="text", text=text, y=paragraph[0].y0))
            paragraph.clear()

    i = 0
    while i < len(lines):
        line = lines[i]
        if i in table_at:
            flush()
            table = table_at[i]
            elements.append(Element(kind="table", rows=table.rows, y=line.y0))
            i = table.end
            continue
        level = is_heading(line, body_size)
        if level:
            flush()
            elements.append(Element(kind="heading", text=re.sub(r"\s+", " ", line.text).strip(), level=level, y=line.y0))
        else:
            if paragraph:
                prev = paragraph[-1]
                gap = line.y0 - prev.y1
                if gap > 0.7 * prev.size or gap < -prev.size or _BULLET.match(line.text):
                    flush()
            paragraph.append(line)
        i += 1
    flush()
    return elements


# ── PDF ──────────────────────────────────────────────────────────────────────
@dataclass
class _Span:
    text: str
    x0: float
    y0: float
    x1: float
    y1: float
    size: float
    bold: bool


def _spans_to_lines(spans: list[_Span]) -> list[Line]:
    """Group spans into visual rows. Wide horizontal gaps become cell separators."""
    lines: list[Line] = []
    rows: list[list[_Span]] = []
    for span in sorted(spans, key=lambda s: ((s.y0 + s.y1) / 2, s.x0)):
        yc = (span.y0 + span.y1) / 2
        if rows:
            last = rows[-1]
            last_yc = sum((s.y0 + s.y1) / 2 for s in last) / len(last)
            if abs(yc - last_yc) <= max(2.5, 0.4 * span.size):
                last.append(span)
                continue
        rows.append([span])
    for row in rows:
        row.sort(key=lambda s: s.x0)
        text = ""
        for idx, span in enumerate(row):
            if idx:
                gap = span.x0 - row[idx - 1].x1
                text += "   " if gap > 12 else (" " if gap > 1 and not text.endswith(" ") else "")
            text += span.text
        text = text.rstrip()
        if text.strip():
            chars = sum(len(s.text) for s in row) or 1
            lines.append(
                Line(
                    text=text,
                    y0=min(s.y0 for s in row),
                    y1=max(s.y1 for s in row),
                    size=max(s.size for s in row),
                    bold=sum(len(s.text) for s in row if s.bold) / chars > 0.6,
                )
            )
    return lines


def _numeric_share(group: list[_Span]) -> float:
    return sum(is_numeric_cell(s.text.strip()) for s in group) / len(group)


def _gutter(
    spans: list[_Span], fractions: tuple[float, ...], *, max_crossing: float, min_side: int
) -> tuple[list[_Span], list[_Span]] | None:
    """Find a vertical whitespace gutter and return ``(left, right)`` spans, or ``None``.

    Positions are tried relative to the text's own extent, so the same search works
    on a whole page and on one half of a spread. The few spans that cross the gutter
    (full-width headings) go with the left side.
    """
    x_min, x_max = min(s.x0 for s in spans), max(s.x1 for s in spans)
    for fraction in fractions:
        x = x_min + (x_max - x_min) * fraction
        crossing = [s for s in spans if s.x0 < x - 2 and s.x1 > x + 2]
        if len(crossing) > max_crossing * len(spans):
            continue
        left = [s for s in spans if s.x1 <= x + 2]
        right = [s for s in spans if s.x0 >= x - 2]
        if len(left) >= min_side and len(right) >= min_side:
            return left + crossing, right
    return None


def _split_columns(spans: list[_Span], width: float, height: float = 0.0) -> list[list[_Span]]:
    """Return spans grouped into columns, in reading order.

    Handles two layouts found in real filings, and their combination:

    * **Spreads** — a landscape PDF page holding two printed pages side by side.
      These are always split down the middle, whatever they contain; otherwise a
      balance sheet on the left page and a P&L on the right would share rows.
    * **Text columns** — split only when both sides are prose. A table with labels
      on the left and figures on the right also has a gutter and must stay whole.
    """
    if not spans:
        return [spans]
    parts = [spans]
    if height and width > height * 1.2 and len(spans) >= 10:
        halves = _gutter(spans, (0.5, 0.49, 0.51, 0.48, 0.52), max_crossing=0.04, min_side=4)
        if halves is not None:
            parts = list(halves)

    columns: list[list[_Span]] = []
    for part in parts:
        split = None
        if len(part) >= 30:
            split = _gutter(part, (0.5, 0.47, 0.53, 0.44, 0.56, 0.41, 0.59), max_crossing=0.08, min_side=12)
        if split is not None and _numeric_share(split[0]) <= 0.25 and _numeric_share(split[1]) <= 0.25:
            columns.extend(split)
        else:
            columns.append(part)
    return columns


def _page_spans(page, exclude: list[tuple[float, float, float, float]]) -> list[_Span]:  # noqa: ANN001
    spans: list[_Span] = []
    for block in page.get_text("dict").get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                text = span.get("text", "")
                if not text.strip():
                    continue
                x0, y0, x1, y1 = span["bbox"]
                xc, yc = (x0 + x1) / 2, (y0 + y1) / 2
                if any(bx0 <= xc <= bx1 and by0 <= yc <= by1 for bx0, by0, bx1, by1 in exclude):
                    continue
                spans.append(_Span(text, x0, y0, x1, y1, float(span.get("size", 10)), bool(span.get("flags", 0) & 16)))
    return spans


def _ruled_table_pages(doc) -> list[int]:  # noqa: ANN001
    """Pages that could contain a fully ruled table.

    Ruled-table detection is by far the slowest step, and most pages of a designed
    annual report have no grid at all. A page qualifies only if it draws several
    horizontal *and* vertical rules (or cell rectangles) and carries enough figures.
    """
    candidates: list[int] = []
    for index in range(min(doc.page_count, MAX_PAGES)):
        page = doc[index]
        # A grid has rules at several distinct positions in both directions; decorative
        # lines, underlines and page borders do not.
        rows_at: set[int] = set()
        cols_at: set[int] = set()
        try:
            drawings = page.get_cdrawings()
        except Exception:
            drawings = []
        for drawing in drawings:
            x0, y0, x1, y1 = drawing.get("rect", (0, 0, 0, 0))
            w, h = abs(x1 - x0), abs(y1 - y0)
            if h <= 2 and w >= 30:
                rows_at.add(round(y0 / 3))
            elif w <= 2 and h >= 12:
                cols_at.add(round(x0 / 3))
            elif 10 <= h <= 60 and 20 <= w <= 400:   # a drawn cell
                rows_at.update((round(y0 / 3), round(y1 / 3)))
                cols_at.update((round(x0 / 3), round(x1 / 3)))
        if len(rows_at) < 4 or len(cols_at) < 3:
            continue
        if sum(is_numeric_cell(word[4]) for word in page.get_text("words")) >= 8:
            candidates.append(index)
    return candidates


def _plumber_tables(
    path: Path, warnings: list[str], pages: list[int] | None = None
) -> dict[int, list[tuple[tuple[float, float, float, float], list[list[str]]]]]:
    """Ruled tables per page index, as ``(bbox, rows)``. ``pages`` limits the scan."""
    found: dict[int, list] = {}
    if pages is not None and not pages:
        return found
    try:
        import pdfplumber

        started = time.monotonic()
        with pdfplumber.open(str(path)) as pdf:
            indices = pages if pages is not None else range(min(len(pdf.pages), MAX_PAGES))
            for index in indices:
                if time.monotonic() - started > RULED_TABLE_BUDGET_SECONDS:
                    # Unruled detection still covers the remaining pages.
                    warnings.append("Ruled-table detection was cut short on this large PDF; remaining tables were read from text.")
                    break
                page = pdf.pages[index]
                try:
                    tables = page.find_tables()
                except Exception:
                    continue
                for table in tables:
                    rows = [
                        [re.sub(r"\s+", " ", c).strip() if c else "" for c in row]
                        for row in (table.extract() or [])
                    ]
                    rows = [r for r in rows if any(r)]
                    cells = [c for r in rows for c in r if c]
                    if len(rows) < 2 or max((len(r) for r in rows), default=0) < 2 or not cells:
                        continue
                    # Decorative boxes and layout frames are not data tables.
                    if sum(is_numeric_cell(c) for c in cells) / len(cells) < 0.15:
                        continue
                    found.setdefault(index, []).append((tuple(table.bbox), rows))
    except Exception as exc:
        warnings.append("Ruled-table detection was unavailable for this PDF; tables were read from text.")
        logger.warning("pdfplumber failed", extra={"error": type(exc).__name__})
    return found


def _ocr_lines(page) -> list[Line]:  # noqa: ANN001
    """OCR a page through PyMuPDF's Tesseract bridge. Raises if Tesseract is missing."""
    textpage = page.get_textpage_ocr(dpi=200, full=True)
    text = page.get_text("text", textpage=textpage)
    return [Line(text=t, y0=i * 14.0, y1=i * 14.0 + 10) for i, t in enumerate(text.splitlines()) if t.strip()]


def _strip_running_lines(pages: list[tuple[list[Line], float]]) -> None:
    """Drop page numbers and repeated headers/footers (same text in the margins of most pages)."""
    def key(text: str) -> str:
        return re.sub(r"\d+", "#", text.lower()).strip()

    def in_margin(line: Line, height: float) -> bool:
        return line.y1 < height * 0.08 or line.y0 > height * 0.92

    counts: Counter[str] = Counter()
    for lines, height in pages:
        counts.update({key(ln.text) for ln in lines if in_margin(ln, height)})
    threshold = max(2, round(len(pages) * 0.6))
    repeated = {k for k, c in counts.items() if c >= threshold}
    for lines, height in pages:
        lines[:] = [
            ln for ln in lines
            if not (in_margin(ln, height) and (key(ln.text) in repeated or _PAGE_NUMBER.match(ln.text.strip())))
        ]


def parse_pdf(path: Path) -> ParsedDocument:
    import pymupdf

    try:
        doc = pymupdf.open(str(path))
    except Exception as exc:
        raise InvalidDocumentError("The file could not be opened as a PDF. It may be corrupted.") from exc
    try:
        if doc.needs_pass:
            raise InvalidDocumentError("The PDF is password-protected. Remove the password and upload again.")
        if doc.page_count == 0:
            raise EmptyDocumentError("The PDF contains no pages.")

        parsed = ParsedDocument(file_type="pdf")
        if doc.page_count > MAX_PAGES:
            parsed.warnings.append(f"Only the first {MAX_PAGES} pages were processed.")
        tables = _plumber_tables(path, parsed.warnings, _ruled_table_pages(doc))

        raw: list[tuple[list[Line], float]] = []
        image_only = ocr_ok = ocr_failed = 0
        for index in range(min(doc.page_count, MAX_PAGES)):
            page = doc[index]
            page_tables = tables.get(index, [])
            spans = _page_spans(page, [bbox for bbox, _ in page_tables])
            lines: list[Line] = []
            for column in _split_columns(spans, page.rect.width, page.rect.height):
                lines.extend(_spans_to_lines(column))
            if not lines and not page_tables and page.get_images(full=False):
                image_only += 1
                try:
                    lines = _ocr_lines(page)
                    ocr_ok += 1
                except Exception:
                    ocr_failed += 1
            raw.append((lines, page.rect.height))

        _strip_running_lines(raw)

        sizes = [ln.size for lines, _ in raw for ln in lines for _ in range(max(1, len(ln.text) // 20))]
        body_size = statistics.median(sizes) if sizes else 10.0

        for index, (lines, _height) in enumerate(raw):
            elements = build_elements(lines, body_size)
            for bbox, rows in tables.get(index, []):
                elements.append(Element(kind="table", rows=rows, y=bbox[1]))
            elements.sort(key=lambda el: el.y)
            parsed.pages.append(ParsedPage(number=index + 1, elements=elements))

        if ocr_ok:
            parsed.warnings.append(f"{ocr_ok} scanned page(s) were read with OCR; accuracy may be lower.")
        if ocr_failed:
            parsed.warnings.append(
                f"{ocr_failed} scanned page(s) could not be read because OCR is unavailable or failed."
            )
        if parsed.char_count < 20:
            if image_only:
                raise OCRError(
                    "This PDF appears to be scanned images and OCR could not extract text. "
                    "Install Tesseract or upload a text-based PDF."
                )
            raise EmptyDocumentError("No text could be extracted from the PDF.")
        return parsed
    finally:
        doc.close()


# ── DOCX ─────────────────────────────────────────────────────────────────────
def parse_docx(path: Path) -> ParsedDocument:
    try:
        from docx import Document as load_docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = load_docx(str(path))
    except Exception as exc:
        raise InvalidDocumentError("The file could not be opened as a Word document.") from exc

    parsed = ParsedDocument(file_type="docx")
    page = ParsedPage(number=1)
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            xml = child.xml
            # Honour explicit and rendered page breaks so citations carry real page numbers.
            if ('w:type="page"' in xml or "lastRenderedPageBreak" in xml) and page.elements:
                parsed.pages.append(page)
                page = ParsedPage(number=page.number + 1)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            if style.lower().startswith(("heading", "title")):
                level = 1 if style.lower() in ("title", "heading 1") else 2
                page.elements.append(Element(kind="heading", text=text, level=level))
            elif is_heading(Line(text=text, y0=0, y1=0), 10.0):
                page.elements.append(Element(kind="heading", text=text))
            else:
                page.elements.append(Element(kind="text", text=text))
        elif tag == "tbl":
            rows = []
            for row in Table(child, doc).rows:
                cells, previous = [], None
                for cell in row.cells:
                    if cell._tc is previous:  # merged cells repeat the same underlying cell
                        continue
                    previous = cell._tc
                    cells.append(re.sub(r"\s+", " ", cell.text).strip())
                if any(cells):
                    rows.append(cells)
            if rows:
                page.elements.append(Element(kind="table", rows=rows))
    parsed.pages.append(page)
    if parsed.char_count < 20:
        raise EmptyDocumentError("The Word document contains no readable text.")
    return parsed


# ── XLSX ─────────────────────────────────────────────────────────────────────
def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:,.2f}".rstrip("0").rstrip(".") if not value.is_integer() else f"{int(value):,}"
    if isinstance(value, int) and not isinstance(value, bool):
        # Bare years must stay unformatted so they are recognised as periods.
        return str(value) if 1900 <= value <= 2100 else f"{value:,}"
    if hasattr(value, "strftime"):
        return value.strftime("%B %d, %Y")
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_xlsx(path: Path) -> ParsedDocument:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise InvalidDocumentError("The file could not be opened as an Excel workbook.") from exc

    parsed = ParsedDocument(file_type="xlsx")
    try:
        for number, sheet in enumerate(workbook.worksheets, start=1):
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [_cell_text(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    rows.append(cells)
                if len(rows) >= 5000:
                    parsed.warnings.append(f"Sheet '{sheet.title}' was truncated to 5,000 rows.")
                    break
            page = ParsedPage(number=number, elements=[Element(kind="heading", text=sheet.title, level=1)])
            # Split the sheet into blocks at single-cell caption rows so each statement is its own table.
            block: list[list[str]] = []
            for cells in rows:
                filled = [c for c in cells if c]
                if len(filled) == 1 and not is_numeric_cell(filled[0]) and len(block) >= 2:
                    page.elements.append(Element(kind="table", rows=block))
                    block = []
                    page.elements.append(Element(kind="text", text=filled[0]))
                elif len(filled) == 1 and not block:
                    page.elements.append(Element(kind="text", text=filled[0]))
                else:
                    block.append(cells)
            if block:
                page.elements.append(Element(kind="table", rows=block))
            parsed.pages.append(page)
    finally:
        workbook.close()
    if parsed.char_count < 20:
        raise EmptyDocumentError("The workbook contains no data.")
    return parsed


# ── Plain text / Markdown ────────────────────────────────────────────────────
def parse_text(path: Path) -> ParsedDocument:
    raw = path.read_bytes()
    try:
        content = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        content = raw.decode("latin-1")
    if len(content.strip()) < 20:
        raise EmptyDocumentError("The text file is empty.")

    parsed = ParsedDocument(file_type="text")
    for number, page_text in enumerate(content.split("\f"), start=1):
        lines: list[Line] = []
        y = 0.0
        source = page_text.splitlines()
        for idx, text in enumerate(source):
            if not text.strip():
                y += 12.0  # blank line → paragraph gap
                continue
            stripped = text.strip()
            if re.fullmatch(r"[|:\- ]{3,}", stripped) or re.fullmatch(r"[=\-]{3,}", stripped):
                continue  # markdown table separators / setext underlines
            forced = 0
            md = re.match(r"^(#{1,6})\s+(.*)$", stripped)
            if md:
                forced, stripped = (1 if len(md.group(1)) == 1 else 2), md.group(2).strip()
            elif idx + 1 < len(source) and re.fullmatch(r"[=\-]{3,}", source[idx + 1].strip()):
                forced = 1
            if stripped.startswith("|") and stripped.endswith("|"):
                stripped = "   ".join(c.strip() for c in stripped.strip("|").split("|"))
            lines.append(Line(text=stripped, y0=y, y1=y + 10.0, forced_heading=forced))
            y += 11.0
        elements = build_elements(lines, 10.0)
        if elements:
            parsed.pages.append(ParsedPage(number=number, elements=elements))
    if not parsed.pages:
        raise EmptyDocumentError("The text file contains no readable content.")
    for index, page in enumerate(parsed.pages, start=1):
        page.number = index
    return parsed


_PARSERS = {".pdf": parse_pdf, ".docx": parse_docx, ".xlsx": parse_xlsx, ".txt": parse_text, ".md": parse_text}


def parse_document(path: Path, extension: str | None = None) -> ParsedDocument:
    ext = (extension or path.suffix).lower()
    parser = _PARSERS.get(ext)
    if parser is None:
        raise UnsupportedFormatError(f"Unsupported file type '{ext}'.")
    return parser(path)
